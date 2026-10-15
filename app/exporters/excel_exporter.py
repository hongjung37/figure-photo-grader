from __future__ import annotations

import argparse
import json
import math
import os
import sqlite3
import sys
import tempfile
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.config import DEFAULT_DB_PATH
from app.db import ensure_tag_tables, sync_analysis_tags
from app.models.analysis import (
    AESTHETIC_TAG_CATEGORIES,
    AESTHETIC_TAG_OPTIONS,
    PhotoAnalysis,
)

SHEET_NAME = "사진 분석"
HELP_SHEET_NAME = "사용 안내"
ID_HEADER = "ID"
STATUS_HEADER = "검토 상태"
USER_NOTES_HEADER = "사용자 메모"
REVIEW_STATUSES = ("미검토", "확인 필요", "검토 완료", "토너먼트 1위")

BASE_COLUMNS = (
    ID_HEADER,
    "분석 시각(UTC)",
    "파일명",
    "SHA-256",
    "너비(px)",
    "높이(px)",
    "형식",
    "모델",
    "촬영 정보(JSON)",
)
ANALYSIS_HEADERS = {
    "구도 점수": "composition_score",
    "구도 메모": "composition_notes",
    "조명 점수": "lighting_score",
    "조명 메모": "lighting_notes",
    "색감 점수": "color_score",
    "색감 메모": "color_notes",
    "촬영 기술 점수": "technique_score",
    "촬영 기술 메모": "technique_notes",
    "최종 점수": "final_score",
    "최종 평가": "final_reason",
    "강점(줄바꿈 구분)": "strengths",
    "개선점(줄바꿈 구분)": "improvements",
    "피규어 유형": "figure_category",
    "초점 상태": "focus_status",
    "초점 분석": "focus_notes",
    "워터마크 촬영 정보(JSON)": "detected_camera_settings",
}
AESTHETIC_TAG_HEADERS = {
    f"{category_label} 태그(줄바꿈 구분)": category_key
    for category_key, category_label in AESTHETIC_TAG_CATEGORIES.items()
}
EXPORT_COLUMNS = (
    *BASE_COLUMNS,
    *ANALYSIS_HEADERS.keys(),
    *AESTHETIC_TAG_HEADERS.keys(),
    STATUS_HEADER,
    USER_NOTES_HEADER,
)
REQUIRED_EXPORT_COLUMNS = (*BASE_COLUMNS, *ANALYSIS_HEADERS.keys(), STATUS_HEADER, USER_NOTES_HEADER)
SCORE_HEADERS = frozenset(
    header
    for header, field in ANALYSIS_HEADERS.items()
    if field.endswith("_score")
)
LIST_HEADERS = frozenset(
    header
    for header, field in ANALYSIS_HEADERS.items()
    if field in {"strengths", "improvements"}
)
TEXT_HEADERS = frozenset(EXPORT_COLUMNS) - SCORE_HEADERS - {ID_HEADER, "너비(px)", "높이(px)"}
MIGRATIONS = {
    "review_status": "TEXT NOT NULL DEFAULT '미검토'",
    "user_notes": "TEXT NOT NULL DEFAULT ''",
}


class ExcelManagerError(ValueError):
    """A workbook or database cannot be safely synchronized."""


def _require_xlsx(path: Path, description: str) -> None:
    if path.suffix.lower() != ".xlsx":
        raise ExcelManagerError(f"{description} 파일은 .xlsx 형식이어야 합니다: {path}")


def _open_database(db_path: Path, *, read_only: bool = False) -> sqlite3.Connection:
    if not db_path.is_file():
        raise ExcelManagerError(f"SQLite DB 파일을 찾을 수 없습니다: {db_path}")
    if read_only:
        connection = sqlite3.connect(db_path.resolve().as_uri() + "?mode=ro", uri=True)
    else:
        connection = sqlite3.connect(db_path)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    return connection


def _validate_schema(connection: sqlite3.Connection) -> set[str]:
    table = connection.execute(
        "SELECT name FROM sqlite_master WHERE type = 'table' AND name = 'photo_analyses'"
    ).fetchone()
    if table is None:
        raise ExcelManagerError("DB에 photo_analyses 테이블이 없습니다. 먼저 사진을 분석해 저장해 주세요.")

    columns = {row["name"] for row in connection.execute("PRAGMA table_info(photo_analyses)")}
    required = {
        "id",
        "sha256",
        "source_filename",
        "width",
        "height",
        "image_format",
        "model",
        "exif_json",
        "analysis_json",
        "created_at",
    }
    missing = required - columns
    if missing:
        raise ExcelManagerError(f"DB 테이블의 필수 열이 없습니다: {', '.join(sorted(missing))}")
    return columns


def _ensure_schema(connection: sqlite3.Connection) -> None:
    columns = _validate_schema(connection)
    # Add only two non-destructive columns; existing analyses receive safe defaults.
    for name, definition in MIGRATIONS.items():
        if name not in columns:
            connection.execute(f"ALTER TABLE photo_analyses ADD COLUMN {name} {definition}")
    ensure_tag_tables(connection)


def _get_export_rows(db_path: Path) -> list[dict[str, Any]]:
    connection = _open_database(db_path, read_only=True)
    try:
        _validate_schema(connection)
        return [dict(row) for row in connection.execute("SELECT * FROM photo_analyses ORDER BY id")]
    finally:
        connection.close()


def _excel_record(row: dict[str, Any]) -> dict[str, Any]:
    try:
        analysis = PhotoAnalysis.model_validate_json(row["analysis_json"])
        exif = json.loads(row["exif_json"])
    except (ValueError, TypeError, KeyError) as exc:
        raise ExcelManagerError(f"DB 레코드 ID {row.get('id', '?')}의 JSON 데이터가 올바르지 않습니다: {exc}") from exc
    if not isinstance(exif, dict):
        raise ExcelManagerError(f"DB 레코드 ID {row['id']}의 촬영 정보가 JSON 객체가 아닙니다.")

    record: dict[str, Any] = {
        ID_HEADER: row["id"],
        "분석 시각(UTC)": row["created_at"],
        "파일명": row["source_filename"],
        "SHA-256": row["sha256"],
        "너비(px)": row["width"],
        "높이(px)": row["height"],
        "형식": row["image_format"],
        "모델": row["model"],
        "촬영 정보(JSON)": json.dumps(exif, ensure_ascii=False, sort_keys=True),
    }
    analysis_data = analysis.model_dump(mode="json")
    for header, field in ANALYSIS_HEADERS.items():
        value = analysis_data[field]
        if isinstance(value, dict):
            record[header] = json.dumps(value, ensure_ascii=False) if value else ""
        elif isinstance(value, list):
            record[header] = "\n".join(value)
        elif value is None:
            record[header] = ""
        else:
            record[header] = value
    tag_data = analysis.aesthetic_tags.model_dump(mode="python")
    for header, category_key in AESTHETIC_TAG_HEADERS.items():
        record[header] = "\n".join(tag_data[category_key])
    record[STATUS_HEADER] = row.get("review_status") or REVIEW_STATUSES[0]
    record[USER_NOTES_HEADER] = row.get("user_notes") or ""
    return record


def _format_workbook(path: Path, row_count: int) -> None:
    workbook = load_workbook(path)
    sheet = workbook[SHEET_NAME]
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 34

    header_fill = PatternFill("solid", fgColor="183153")
    for cell in sheet[1]:
        cell.font = Font(name="맑은 고딕", color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.comment = None

    header_positions = {cell.value: cell.column for cell in sheet[1]}
    preferred_widths = {
        ID_HEADER: 9,
        "분석 시각(UTC)": 24,
        "파일명": 28,
        "SHA-256": 24,
        "너비(px)": 12,
        "높이(px)": 12,
        "형식": 11,
        "모델": 24,
        "촬영 정보(JSON)": 38,
        "구도 점수": 12,
        "조명 점수": 12,
        "색감 점수": 12,
        "촬영 기술 점수": 15,
        "최종 점수": 12,
        "피규어 유형": 18,
        "초점 상태": 14,
        "초점 분석": 30,
        "워터마크 촬영 정보(JSON)": 32,
        STATUS_HEADER: 14,
        USER_NOTES_HEADER: 36,
    }
    preferred_widths.update({header: 24 for header in AESTHETIC_TAG_HEADERS})
    long_text_headers = {
        "촬영 정보(JSON)",
        "구도 메모",
        "조명 메모",
        "색감 메모",
        "촬영 기술 메모",
        "최종 평가",
        "강점(줄바꿈 구분)",
        "개선점(줄바꿈 구분)",
        USER_NOTES_HEADER,
    } | set(AESTHETIC_TAG_HEADERS)
    for header, column in header_positions.items():
        cells = [sheet.cell(row, column) for row in range(2, row_count + 2)]
        width = preferred_widths.get(header)
        if width is None:
            width = min(max((len(str(cell.value or "")) for cell in cells), default=12) + 2, 42)
        sheet.column_dimensions[sheet.cell(1, column).column_letter].width = width
        for cell in cells:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=header in long_text_headers,
            )
            # Force all exported user-controlled text to literal text, not an Excel formula.
            if header in TEXT_HEADERS and cell.value is not None:
                cell.data_type = "s"

    if row_count:
        table = Table(displayName="PhotoAnalyses", ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)

        status_column = header_positions[STATUS_HEADER]
        status_list_str = ",".join(REVIEW_STATUSES)
        validation = DataValidation(
            type="list",
            formula1=f'"{status_list_str}"',
            allow_blank=False,
        )
        validation.error = "목록에서 검토 상태를 선택해 주세요."
        validation.errorTitle = "검토 상태 확인"
        validation.prompt = f"{' / '.join(REVIEW_STATUSES)} 중에서 선택합니다."
        validation.promptTitle = "검토 상태"
        sheet.add_data_validation(validation)
        validation.add(f"{sheet.cell(2, status_column).coordinate}:{sheet.cell(max(row_count + 1, 1000), status_column).coordinate}")

        for header in SCORE_HEADERS:
            column = header_positions[header]
            letter = sheet.cell(1, column).column_letter
            sheet.conditional_formatting.add(
                f"{letter}2:{letter}{row_count + 1}",
                ColorScaleRule(
                    start_type="num",
                    start_value=0,
                    start_color="F8696B",
                    mid_type="num",
                    mid_value=5,
                    mid_color="FFEB84",
                    end_type="num",
                    end_value=10,
                    end_color="63BE7B",
                ),
            )

    instructions = [
        ("사진 분석 엑셀 관리", "설명"),
        ("점수 기준", "9.0~10: 프로사진작가 및 전문가급 | 8.0~8.9: 매우 우수 | 7.0~7.9: 우수 | 5.0~6.9: 보통 (일반 스냅) | 5.0 미만: 개선 필요"),
        ("편집 가능", "점수, 평가 메모([장점]/[단점] 구분), 강점·개선점, 검토 상태, 사용자 메모"),
        ("태그 편집", "범주별 태그 열에서 줄바꿈으로 구분합니다. 허용 어휘만 사용할 수 있고 범주당 최대 3개입니다."),
        ("편집 제한", "ID 및 파일·모델·촬영 정보는 표시용이며 SQLite로 가져올 때 변경되지 않습니다."),
        ("반영 방법", "저장한 뒤 photo_excel_manager.py import 명령으로 가져옵니다."),
        ("안전 장치", "가져오기는 ID로 기존 행만 갱신합니다. 새 행 추가·행 삭제는 DB에 반영하지 않습니다."),
        ("검증", "중복/미등록 ID, 잘못된 점수(0~10), 누락 열 또는 잘못된 상태가 있으면 전체 가져오기가 취소됩니다."),
        ("검토 상태", " / ".join(REVIEW_STATUSES)),
        ("개인정보", "촬영 정보에는 DB에 저장된 허용 EXIF만 포함되며 GPS·촬영 시각 등 제외 정보는 내보내지 않습니다."),
        ("공유 주의", "파일명, 평가 및 촬영 설정이 포함될 수 있으니 엑셀 파일을 안전하게 보관하세요."),
    ]
    instructions.extend(
        (
            f"{AESTHETIC_TAG_CATEGORIES[category_key]} 태그 어휘",
            " · ".join(options),
        )
        for category_key, options in AESTHETIC_TAG_OPTIONS.items()
    )
    help_sheet = workbook.create_sheet(HELP_SHEET_NAME)
    help_sheet.sheet_view.showGridLines = False
    help_sheet.append(("항목", "설명"))
    for item in instructions[1:]:
        help_sheet.append(item)
    for cell in help_sheet[1]:
        cell.font = Font(name="맑은 고딕", color="FFFFFF", bold=True)
        cell.fill = header_fill
    help_sheet.column_dimensions["A"].width = 20
    help_sheet.column_dimensions["B"].width = 100
    for row in help_sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        help_sheet.row_dimensions[row[0].row].height = 34
    help_sheet.freeze_panes = "A2"
    workbook.active = 0
    workbook.save(path)


def export_to_excel(db_path: str | Path, output_path: str | Path) -> int:
    database = Path(db_path).expanduser()
    output = Path(output_path).expanduser()
    _require_xlsx(output, "내보내기")
    if database.resolve() == output.resolve():
        raise ExcelManagerError("엑셀 출력 경로가 SQLite DB 경로와 같을 수 없습니다.")

    rows = _get_export_rows(database)
    records = [_excel_record(row) for row in rows]
    frame = pd.DataFrame(records, columns=EXPORT_COLUMNS)
    output.parent.mkdir(parents=True, exist_ok=True)

    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{output.stem}.",
            suffix=".xlsx",
            dir=output.parent,
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)
        with pd.ExcelWriter(temporary_path, engine="openpyxl") as writer:
            frame.to_excel(writer, sheet_name=SHEET_NAME, index=False)
        _format_workbook(temporary_path, len(records))
        os.replace(temporary_path, output)
    except Exception:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)
        raise
    return len(records)


def _is_blank(value: Any) -> bool:
    if value is None:
        return True
    try:
        return bool(pd.isna(value))
    except (TypeError, ValueError):
        return False


def _parse_record_id(value: Any, row_number: int) -> int:
    if _is_blank(value) or isinstance(value, bool):
        raise ExcelManagerError(f"엑셀 {row_number}행의 ID가 비어 있거나 올바르지 않습니다.")
    try:
        number = float(value)
    except (TypeError, ValueError, OverflowError) as exc:
        raise ExcelManagerError(f"엑셀 {row_number}행의 ID는 정수여야 합니다.") from exc
    if not math.isfinite(number) or not number.is_integer() or number <= 0:
        raise ExcelManagerError(f"엑셀 {row_number}행의 ID는 양의 정수여야 합니다.")
    return int(number)


def _required_text(value: Any, header: str, row_number: int) -> str:
    if _is_blank(value):
        return ""
    if not isinstance(value, str):
        raise ExcelManagerError(f"엑셀 {row_number}행의 '{header}' 값은 텍스트여야 합니다.")
    return value


def _parse_analysis_record(record: pd.Series, row_number: int) -> PhotoAnalysis:
    payload: dict[str, Any] = {}
    for header, field in ANALYSIS_HEADERS.items():
        value = record[header]
        if header in SCORE_HEADERS:
            if _is_blank(value) or isinstance(value, bool):
                raise ExcelManagerError(f"엑셀 {row_number}행의 '{header}' 점수가 비어 있습니다.")
            try:
                payload[field] = float(value)
            except (TypeError, ValueError, OverflowError) as exc:
                raise ExcelManagerError(f"엑셀 {row_number}행의 '{header}'는 숫자여야 합니다.") from exc
        elif header in LIST_HEADERS:
            text = _required_text(value, header, row_number)
            payload[field] = [line.strip() for line in text.splitlines() if line.strip()]
        elif field == "detected_camera_settings":
            if _is_blank(value) or not str(value).strip():
                payload[field] = None
            else:
                try:
                    parsed_json = json.loads(str(value).strip())
                    payload[field] = parsed_json if isinstance(parsed_json, dict) else None
                except Exception:
                    payload[field] = None
        else:
            payload[field] = _required_text(value, header, row_number)

    tag_data: dict[str, list[str]] = {}
    for header, category_key in AESTHETIC_TAG_HEADERS.items():
        value = record[header] if header in record.index else ""
        text = _required_text(value, header, row_number)
        tag_data[category_key] = [line.strip() for line in text.splitlines() if line.strip()]
    payload["aesthetic_tags"] = tag_data

    try:
        return PhotoAnalysis.model_validate(payload)
    except Exception as exc:
        raise ExcelManagerError(f"엑셀 {row_number}행의 분석 내용이 유효하지 않습니다: {exc}") from exc


def import_from_excel(db_path: str | Path, input_path: str | Path) -> int:
    database = Path(db_path).expanduser()
    workbook_path = Path(input_path).expanduser()
    _require_xlsx(workbook_path, "가져오기")
    if not workbook_path.is_file():
        raise ExcelManagerError(f"엑셀 파일을 찾을 수 없습니다: {workbook_path}")
    if database.resolve() == workbook_path.resolve():
        raise ExcelManagerError("엑셀 입력 경로가 SQLite DB 경로와 같을 수 없습니다.")

    try:
        frame = pd.read_excel(workbook_path, sheet_name=SHEET_NAME, dtype={ID_HEADER: "object"})
    except (ValueError, OSError, ImportError) as exc:
        raise ExcelManagerError(f"엑셀 파일을 읽지 못했습니다: {exc}") from exc

    missing_headers = set(REQUIRED_EXPORT_COLUMNS) - set(frame.columns)
    if missing_headers:
        raise ExcelManagerError(f"엑셀에 필요한 열이 없습니다: {', '.join(sorted(missing_headers))}")

    parsed_rows: list[tuple[int, PhotoAnalysis, str, str]] = []
    seen_ids: set[int] = set()
    for index, record in frame.iterrows():
        row_number = int(index) + 2
        record_id = _parse_record_id(record[ID_HEADER], row_number)
        if record_id in seen_ids:
            raise ExcelManagerError(f"엑셀에 중복된 ID가 있습니다: {record_id}")
        seen_ids.add(record_id)

        analysis = _parse_analysis_record(record, row_number)
        status = _required_text(record[STATUS_HEADER], STATUS_HEADER, row_number)
        if status not in REVIEW_STATUSES:
            raise ExcelManagerError(
                f"엑셀 {row_number}행의 검토 상태는 다음 중 하나여야 합니다: {', '.join(REVIEW_STATUSES)}"
            )
        user_notes = _required_text(record[USER_NOTES_HEADER], USER_NOTES_HEADER, row_number)
        parsed_rows.append((record_id, analysis, status, user_notes))

    connection = _open_database(database)
    try:
        connection.execute("BEGIN IMMEDIATE")
        _ensure_schema(connection)
        existing_ids = {
            row["id"]
            for row in connection.execute("SELECT id FROM photo_analyses")
        }
        unknown_ids = seen_ids - existing_ids
        if unknown_ids:
            raise ExcelManagerError(f"DB에 없는 ID는 가져올 수 없습니다: {', '.join(map(str, sorted(unknown_ids)))}")

        for record_id, analysis, status, user_notes in parsed_rows:
            connection.execute(
                "UPDATE photo_analyses SET analysis_json = ?, review_status = ?, user_notes = ? WHERE id = ?",
                (analysis.model_dump_json(), status, user_notes, record_id),
            )
            sync_analysis_tags(connection, record_id, analysis.aesthetic_tags)
        connection.commit()
        return len(parsed_rows)
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


def summarize_tags(db_path: str | Path) -> dict[str, Any]:
    """Return read-only counts of tagged records and each controlled tag."""
    connection = _open_database(Path(db_path).expanduser(), read_only=True)
    try:
        _validate_schema(connection)
        tables = {
            row["name"]
            for row in connection.execute("SELECT name FROM sqlite_master WHERE type = 'table'")
        }
        required_tables = {"photo_tag_catalog", "photo_analysis_tags"}
        if not required_tables <= tables:
            raise ExcelManagerError("DB에 태그 테이블이 없습니다. 먼저 새 분석을 저장하거나 태그 Excel을 가져와 주세요.")

        total_records = int(connection.execute("SELECT COUNT(*) FROM photo_analyses").fetchone()[0])
        tagged_records = int(
            connection.execute(
                "SELECT COUNT(DISTINCT analysis_id) FROM photo_analysis_tags"
            ).fetchone()[0]
        )
        usage_rows = connection.execute(
            """SELECT c.category_key, c.category_label, c.tag, c.sort_order,
                      COUNT(DISTINCT t.analysis_id) AS usage_count
               FROM photo_tag_catalog AS c
               LEFT JOIN photo_analysis_tags AS t
                 ON t.category_key = c.category_key AND t.tag = c.tag
               GROUP BY c.category_key, c.category_label, c.tag, c.sort_order
               ORDER BY c.category_key, c.sort_order"""
        ).fetchall()
        categories: dict[str, list[dict[str, Any]]] = {
            label: [] for label in AESTHETIC_TAG_CATEGORIES.values()
        }
        for row in usage_rows:
            if row["usage_count"]:
                categories[row["category_label"]].append(
                    {"tag": row["tag"], "count": int(row["usage_count"])}
                )
        return {
            "total_records": total_records,
            "tagged_records": tagged_records,
            "untagged_records": max(0, total_records - tagged_records),
            "categories": categories,
        }
    finally:
        connection.close()


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="사진 분석 결과를 Excel로 내보내고, 검증 후 SQLite에 반영합니다."
    )
    subparsers = parser.add_subparsers(dest="command", required=True)

    export_parser = subparsers.add_parser("export", help="SQLite 분석 내용을 Excel로 내보냅니다.")
    export_parser.add_argument("--db", type=Path, default=DEFAULT_DB_PATH, help="SQLite DB 경로")
    export_parser.add_argument(
        "--output",
        type=Path,
        help="출력 .xlsx 경로 (기본값: DB 폴더의 photo_analyses.xlsx)",
    )

    import_parser = subparsers.add_parser("import", help="편집한 Excel 내용을 SQLite에 반영합니다.")
    import_parser.add_argument("input", type=Path, help="가져올 .xlsx 파일")
    import_parser.add_argument("--db", type=Path, default=DEFAULT_DB_PATH, help="SQLite DB 경로")

    tags_parser = subparsers.add_parser("tags", help="DB의 범주별 태그 사용 빈도를 확인합니다.")
    tags_parser.add_argument("--db", type=Path, default=DEFAULT_DB_PATH, help="SQLite DB 경로")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        if args.command == "export":
            output = args.output or args.db.with_name("photo_analyses.xlsx")
            count = export_to_excel(args.db, output)
            print(f"Excel 내보내기 완료: {output} ({count}건)")
        elif args.command == "import":
            count = import_from_excel(args.db, args.input)
            print(f"Excel 가져오기 완료: {count}건의 기존 레코드를 갱신했습니다. 새 행이나 삭제는 반영하지 않았습니다.")
        else:
            summary = summarize_tags(args.db)
            print(
                f"분석 기록 {summary['total_records']}건 · 태그 있음 {summary['tagged_records']}건 · "
                f"태그 없음 {summary['untagged_records']}건"
            )
            for category_label, items in summary["categories"].items():
                if items:
                    counts = ", ".join(f"{item['tag']} {item['count']}건" for item in items)
                    print(f"{category_label}: {counts}")
        return 0
    except (ExcelManagerError, OSError, sqlite3.Error, ValueError) as exc:
        print(f"오류: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
