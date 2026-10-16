from __future__ import annotations

import json
import sqlite3
import sys
from contextlib import closing
import tempfile
import unittest
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import load_workbook

from photo_analyzer import PhotoAnalysis, PreparedPhoto, save_analysis
from photo_tag_schema import AestheticTags
from photo_excel_manager import (
    ANALYSIS_HEADERS,
    AESTHETIC_TAG_HEADERS,
    ID_HEADER,
    REVIEW_STATUSES,
    SHEET_NAME,
    STATUS_HEADER,
    USER_NOTES_HEADER,
    ExcelManagerError,
    export_to_excel,
    import_from_excel,
    summarize_tags,
)


class PhotoExcelManagerTests(unittest.TestCase):
    def _create_database(self, directory: Path) -> tuple[Path, int]:
        database = directory / "analysis.sqlite3"
        photo = PreparedPhoto(
            source_filename="sample.jpg",
            sha256="a" * 64,
            width=640,
            height=480,
            image_format="JPEG",
            safe_exif={"Make": "Example Camera", "ISOSpeedRatings": 400},
            upload_bytes=b"test-image",
            upload_mime_type="image/jpeg",
        )
        analysis = PhotoAnalysis(
            composition_score=8,
            composition_notes="피사체가 잘 배치되었습니다.",
            lighting_score=7.5,
            lighting_notes="부드러운 빛입니다.",
            color_score=8,
            color_notes="색 조화가 자연스럽습니다.",
            technique_score=7,
            technique_notes="초점이 안정적입니다.",
            final_score=7.8,
            final_reason="전체 완성도가 좋습니다.",
            strengths=["안정된 구도"],
            improvements=["배경을 단순화해 보세요."],
            aesthetic_tags=AestheticTags.empty().model_copy(
                update={"viewpoint": ["정면"], "mood": ["고요함"]}
            ),
        )
        record_id = save_analysis(database, photo, "test-model", analysis)
        return database, record_id

    def test_export_edit_import_preserves_identity_and_saves_review_fields(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            directory = Path(temporary_directory)
            database, record_id = self._create_database(directory)
            workbook_path = directory / "photo_analyses.xlsx"

            self.assertEqual(export_to_excel(database, workbook_path), 1)
            workbook = load_workbook(workbook_path)
            sheet = workbook[SHEET_NAME]
            headers = {cell.value: cell.column for cell in sheet[1]}
            self.assertEqual(sheet.cell(2, headers[ID_HEADER]).value, record_id)
            self.assertEqual(sheet["A1"].value, ID_HEADER)

            sheet.cell(2, headers["최종 점수"], 9.2)
            sheet.cell(2, headers["최종 평가"], "엑셀에서 수정한 평가")
            sheet.cell(2, headers["강점(줄바꿈 구분)"], "구도 안정\n색감 조화")
            sheet.cell(2, headers[STATUS_HEADER], REVIEW_STATUSES[2])
            sheet.cell(2, headers[USER_NOTES_HEADER], "재촬영 후보로 표시")
            sheet.cell(2, headers["시점 태그(줄바꿈 구분)"], "측면\n후면")
            sheet.cell(2, headers["분위기 태그(줄바꿈 구분)"], "따뜻함")
            workbook.save(workbook_path)

            self.assertEqual(import_from_excel(database, workbook_path), 1)
            with closing(sqlite3.connect(database)) as connection:
                connection.row_factory = sqlite3.Row
                row = connection.execute(
                    "SELECT source_filename, sha256, analysis_json, review_status, user_notes "
                    "FROM photo_analyses WHERE id = ?",
                    (record_id,),
                ).fetchone()

            self.assertEqual(row["source_filename"], "sample.jpg")
            self.assertEqual(row["sha256"], "a" * 64)
            restored = PhotoAnalysis.model_validate_json(row["analysis_json"])
            self.assertEqual(restored.final_score, 9.2)
            self.assertEqual(restored.final_reason, "엑셀에서 수정한 평가")
            self.assertEqual(restored.strengths, ["구도 안정", "색감 조화"])
            self.assertEqual(restored.aesthetic_tags.viewpoint, ["측면", "후면"])
            self.assertEqual(restored.aesthetic_tags.mood, ["따뜻함"])
            self.assertEqual(row["review_status"], REVIEW_STATUSES[2])
            self.assertEqual(row["user_notes"], "재촬영 후보로 표시")
            with closing(sqlite3.connect(database)) as connection:
                normalized = connection.execute(
                    "SELECT category_key, tag FROM photo_analysis_tags WHERE analysis_id = ? ORDER BY category_key, tag",
                    (record_id,),
                ).fetchall()
            self.assertEqual(
                normalized,
                [("mood", "따뜻함"), ("viewpoint", "측면"), ("viewpoint", "후면")],
            )
            summary = summarize_tags(database)
            self.assertEqual(summary["tagged_records"], 1)
            self.assertEqual(summary["untagged_records"], 0)
            self.assertEqual(
                summary["categories"]["시점"],
                [{"tag": "측면", "count": 1}, {"tag": "후면", "count": 1}],
            )

    def test_export_does_not_modify_the_sqlite_schema(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            directory = Path(temporary_directory)
            database, _ = self._create_database(directory)
            with closing(sqlite3.connect(database)) as connection:
                before = [row[1] for row in connection.execute("PRAGMA table_info(photo_analyses)")]

            export_to_excel(database, directory / "photo_analyses.xlsx")

            with closing(sqlite3.connect(database)) as connection:
                after = [row[1] for row in connection.execute("PRAGMA table_info(photo_analyses)")]
            self.assertEqual(after, before)
            self.assertNotIn("review_status", after)
            self.assertNotIn("user_notes", after)

    def test_invalid_score_cancels_import_without_partial_updates(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            directory = Path(temporary_directory)
            database, record_id = self._create_database(directory)
            workbook_path = directory / "photo_analyses.xlsx"
            export_to_excel(database, workbook_path)

            workbook = load_workbook(workbook_path)
            sheet = workbook[SHEET_NAME]
            headers = {cell.value: cell.column for cell in sheet[1]}
            sheet.cell(2, headers["최종 점수"], 10.1)
            workbook.save(workbook_path)

            with self.assertRaises(ExcelManagerError):
                import_from_excel(database, workbook_path)

            with closing(sqlite3.connect(database)) as connection:
                row = connection.execute(
                    "SELECT analysis_json FROM photo_analyses WHERE id = ?",
                    (record_id,),
                ).fetchone()
                columns = {item[1] for item in connection.execute("PRAGMA table_info(photo_analyses)")}
            restored = PhotoAnalysis.model_validate_json(row[0])
            self.assertEqual(restored.final_score, 7.8)
            self.assertNotIn("review_status", columns)
            self.assertNotIn("user_notes", columns)

    def test_unknown_id_cannot_insert_or_modify_a_database_record(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            directory = Path(temporary_directory)
            database, _ = self._create_database(directory)
            workbook_path = directory / "photo_analyses.xlsx"
            export_to_excel(database, workbook_path)

            workbook = load_workbook(workbook_path)
            sheet = workbook[SHEET_NAME]
            headers = {cell.value: cell.column for cell in sheet[1]}
            sheet.cell(2, headers[ID_HEADER], 9999)
            workbook.save(workbook_path)

            with self.assertRaises(ExcelManagerError):
                import_from_excel(database, workbook_path)
            with closing(sqlite3.connect(database)) as connection:
                self.assertEqual(connection.execute("SELECT COUNT(*) FROM photo_analyses").fetchone()[0], 1)

    def test_export_includes_the_expected_analysis_fields(self) -> None:
        self.assertEqual(
            set(ANALYSIS_HEADERS.values()),
            set(PhotoAnalysis.model_fields) - {"aesthetic_tags", "cross_check_opinion"},
        )
        self.assertEqual(set(AESTHETIC_TAG_HEADERS.values()), set(AestheticTags.model_fields))

    def test_old_workbook_without_tag_columns_imports_with_empty_tags(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            directory = Path(temporary_directory)
            database, record_id = self._create_database(directory)
            workbook_path = directory / "legacy.xlsx"
            export_to_excel(database, workbook_path)

            workbook = load_workbook(workbook_path)
            sheet = workbook[SHEET_NAME]
            headers = {cell.value: cell.column for cell in sheet[1]}
            for column in sorted(
                (headers[header] for header in AESTHETIC_TAG_HEADERS), reverse=True
            ):
                sheet.delete_cols(column)
            workbook.save(workbook_path)

            self.assertEqual(import_from_excel(database, workbook_path), 1)
            with closing(sqlite3.connect(database)) as connection:
                row = connection.execute(
                    "SELECT analysis_json FROM photo_analyses WHERE id = ?", (record_id,)
                ).fetchone()
                tag_count = connection.execute(
                    "SELECT COUNT(*) FROM photo_analysis_tags WHERE analysis_id = ?", (record_id,)
                ).fetchone()[0]
            restored = PhotoAnalysis.model_validate_json(row[0])
            self.assertEqual(restored.aesthetic_tags.viewpoint, [])
            self.assertEqual(tag_count, 0)


    def test_tournament_winner_status_roundtrip(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            directory = Path(temporary_directory)
            database, record_id = self._create_database(directory)
            # Simulate column added by dashboard or prior import
            with closing(sqlite3.connect(database)) as connection:
                connection.execute("ALTER TABLE photo_analyses ADD COLUMN review_status TEXT DEFAULT '미검토'")
                connection.execute(
                    "UPDATE photo_analyses SET review_status = '토너먼트 1위' WHERE id = ?",
                    (record_id,),
                )
                connection.commit()

            workbook_path = directory / "tournament.xlsx"
            export_to_excel(database, workbook_path)

            workbook = load_workbook(workbook_path)
            sheet = workbook[SHEET_NAME]
            headers = {cell.value: cell.column for cell in sheet[1]}
            self.assertEqual(sheet.cell(2, headers[STATUS_HEADER]).value, "토너먼트 1위")

            # Re-import should succeed without validation errors
            self.assertEqual(import_from_excel(database, workbook_path), 1)
            with closing(sqlite3.connect(database)) as connection:
                status = connection.execute(
                    "SELECT review_status FROM photo_analyses WHERE id = ?", (record_id,)
                ).fetchone()[0]
            self.assertEqual(status, "토너먼트 1위")

    def test_import_preserves_cross_check_opinion(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            directory = Path(temporary_directory)
            database, record_id = self._create_database(directory)

            # Store an analysis record with cross_check_opinion in SQLite
            with closing(sqlite3.connect(database)) as connection:
                connection.row_factory = sqlite3.Row
                row = connection.execute("SELECT analysis_json FROM photo_analyses WHERE id = ?", (record_id,)).fetchone()
                data = json.loads(row["analysis_json"])
                data["cross_check_opinion"] = {
                    "model": "gpt-6-luna",
                    "score": 8.5,
                    "reason": "우수한 디테일",
                    "strengths": ["초점 우수"],
                    "improvements": ["배경 단순화"],
                }
                connection.execute(
                    "UPDATE photo_analyses SET analysis_json = ? WHERE id = ?",
                    (json.dumps(data), record_id),
                )
                connection.commit()

            workbook_path = directory / "export.xlsx"
            export_to_excel(database, workbook_path)

            # Modify user notes in Excel and re-import
            workbook = load_workbook(workbook_path)
            sheet = workbook[SHEET_NAME]
            headers = {cell.value: cell.column for cell in sheet[1]}
            sheet.cell(2, headers[USER_NOTES_HEADER]).value = "수정된 메모"
            workbook.save(workbook_path)

            import_from_excel(database, workbook_path)

            # Verify cross_check_opinion is preserved
            with closing(sqlite3.connect(database)) as connection:
                connection.row_factory = sqlite3.Row
                row = connection.execute("SELECT analysis_json, user_notes FROM photo_analyses WHERE id = ?", (record_id,)).fetchone()
                restored_data = json.loads(row["analysis_json"])
                self.assertEqual(row["user_notes"], "수정된 메모")
                self.assertIsNotNone(restored_data.get("cross_check_opinion"))
                self.assertEqual(restored_data["cross_check_opinion"]["score"], 8.5)
                self.assertEqual(restored_data["cross_check_opinion"]["model"], "gpt-6-luna")

    def test_import_rejects_malformed_camera_settings_json(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            directory = Path(temporary_directory)
            database, record_id = self._create_database(directory)
            workbook_path = directory / "malformed_json.xlsx"
            export_to_excel(database, workbook_path)

            # Corrupt camera settings JSON in Excel
            workbook = load_workbook(workbook_path)
            sheet = workbook[SHEET_NAME]
            headers = {cell.value: cell.column for cell in sheet[1]}
            sheet.cell(2, headers["워터마크 촬영 정보(JSON)"]).value = "{malformed json:"
            workbook.save(workbook_path)

            with self.assertRaises(ExcelManagerError) as ctx:
                import_from_excel(database, workbook_path)
            self.assertIn("JSON 형식", str(ctx.exception))


if __name__ == "__main__":
    unittest.main()
